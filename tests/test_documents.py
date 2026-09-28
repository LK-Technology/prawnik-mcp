"""Offline tests: templates, validation, export gates, MD/DOCX export, examples."""

from __future__ import annotations

import copy
import importlib.util
import re
import shutil
import subprocess
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from prawnik_mcp.contracts import CitationStatus, ClaimCheck, ResultStatus, ReviewerType, TemporalStatus
from prawnik_mcp.documents import money
from prawnik_mcp.documents.render import (
    DOCUMENT_COMPLETE,
    DOCUMENT_INCOMPLETE,
    SIGNATURE_LINE,
    STATUS_LINE,
    binding_hash,
    render_document,
)
from prawnik_mcp.documents.schema import TemplateSpec
from prawnik_mcp.documents.templates import get_template, list_templates, load_all_raw, load_template
from prawnik_mcp.store import Store

ROOT = Path(__file__).resolve().parents[1]
EXAMPLES_DIR = ROOT / "examples"
TODAY = date(2026, 9, 26)
TEMPLATE_IDS = ["wezwanie_do_zaplaty", "reklamacja_konsumencka", "odstapienie_od_umowy_na_odleglosc"]
ALLOWED_SOURCE_DOCS = {"eli:DU/1964/93", "eli:DU/2014/827", "celex:32011L0083"}

_spec = importlib.util.spec_from_file_location("generate_examples", EXAMPLES_DIR / "generate_examples.py")
gen = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(gen)


@pytest.fixture
def store(tmp_path):
    s = Store(tmp_path / "data")
    yield s
    s.close()


def facts_of(tid: str) -> dict:
    return copy.deepcopy(gen.EXAMPLES[tid]["facts"])


def draft_of(tid: str) -> dict:
    return copy.deepcopy(gen.EXAMPLES[tid]["draft"])


def render(store, tid, facts, draft, report_id, out):
    return render_document(store, tid, facts, draft, report_id, out, today=TODAY)


def docx_text(path: Path) -> str:
    return "\n".join(p.text for p in Document(path).paragraphs)


# --------------------------------------------------------------------------- templates


def test_list_and_get_templates():
    ids = {t["template_id"] for t in list_templates()}
    assert ids == set(TEMPLATE_IDS)
    for tid in TEMPLATE_IDS:
        t = get_template(tid)
        assert t and t["version"] and t["required_facts"] and t["qualifying_questions"]
        assert "letter" not in t
    assert get_template("../etc/passwd") is None
    assert get_template("nie_ma_takiego") is None


@pytest.mark.parametrize("tid", TEMPLATE_IDS)
def test_template_metadata_rules(tid):
    raw = load_all_raw()[tid]
    spec = TemplateSpec.model_validate(raw)
    assert spec.review_history == []
    assert "not reviewed by a lawyer" in spec.status
    assert "get_legal_document" in spec.legal_sources_note
    assert spec.supported_dates and spec.exclusions
    for src in spec.legal_sources:
        assert src.document_id in ALLOWED_SOURCE_DOCS
        assert set(src.model_dump()) == {"document_id", "locator", "purpose"}
    for q in spec.qualifying_questions:
        assert q.why_it_matters
    for f in spec.required_facts:
        assert f.type in {"str", "date", "money", "enum", "list"}


def test_reklamacja_qualifying_questions_cover_traps():
    spec = load_template("reklamacja_konsumencka")
    q = {x.id: x for x in spec.qualifying_questions}
    assert "nie" in q["konsument"].excluding_answers
    assert {"treść cyfrowa", "usługa cyfrowa"} <= set(q["przedmiot"].excluding_answers)
    temporal = q["umowa_przed_zmiana_przepisow"]
    assert not temporal.excluding_answers and "tak" in temporal.flag_answers
    assert "nie przesądza" in temporal.flag_note


def test_odstapienie_exclusion_questions_cite_art_38():
    spec = load_template("odstapienie_od_umowy_na_odleglosc")
    wyl = [q for q in spec.qualifying_questions if q.id.startswith("wyl_")]
    assert len(wyl) >= 8
    for q in wyl:
        assert {"document_id": "eli:DU/2014/827", "locator": "art. 38"} in q.related_sources
        assert "zweryfikuj eli:DU/2014/827 art. 38" in q.exclusion_reason


# --------------------------------------------------------------------------- money


@pytest.mark.parametrize("n,words", [
    (1, "jeden"), (12, "dwanaście"), (22, "dwadzieścia dwa"), (100, "sto"), (215, "dwieście piętnaście"),
    (1000, "tysiąc"), (1001, "tysiąc jeden"), (2000, "dwa tysiące"), (5000, "pięć tysięcy"),
    (12000, "dwanaście tysięcy"), (22000, "dwadzieścia dwa tysiące"),
    (1234567, "milion dwieście trzydzieści cztery tysiące pięćset sześćdziesiąt siedem"),
])
def test_int_to_words(n, words):
    assert money.int_to_words(n) == words


def test_amount_words_and_match():
    d = Decimal("1234.50")
    assert money.amount_in_words(d, "PLN") == "tysiąc dwieście trzydzieści cztery złote 50/100"
    assert money.amount_in_words(Decimal("5"), "PLN") == "pięć złotych 00/100"
    assert money.amount_in_words(Decimal("1"), "PLN") == "jeden złoty 00/100"
    assert money.words_match("Tysiąc dwieście trzydzieści cztery złote, 50/100", d, "PLN")
    assert money.words_match("tysiąc dwieście trzydzieści cztery złote pięćdziesiąt groszy", d, "PLN")
    assert not money.words_match("tysiąc dwieście czterdzieści trzy złote 50/100", d, "PLN")
    assert money.format_amount(Decimal("1234567.5"), "PLN") == "1 234 567,50 PLN"


def test_pl_account_checksum():
    assert money.valid_pl_account(gen.FICTIONAL_ACCOUNT)
    assert money.valid_pl_account("PL" + gen.FICTIONAL_ACCOUNT)
    assert not money.valid_pl_account(gen.FICTIONAL_ACCOUNT.replace("1234", "1235"))
    assert not money.valid_pl_account("12345")


# --------------------------------------------------------------------------- binding hash


def test_binding_hash_stable_and_sensitive():
    f = facts_of("wezwanie_do_zaplaty")
    d = draft_of("wezwanie_do_zaplaty")
    h = binding_hash("wezwanie_do_zaplaty", "0.1.0", f, d)
    reordered = dict(reversed(list(f.items())))
    assert binding_hash("wezwanie_do_zaplaty", "0.1.0", reordered, d) == h
    assert re.fullmatch(r"[0-9a-f]{64}", h)
    f2 = copy.deepcopy(f)
    f2["kwota"]["amount"] = "1234.51"
    assert binding_hash("wezwanie_do_zaplaty", "0.1.0", f2, d) != h
    assert binding_hash("wezwanie_do_zaplaty", "0.1.1", f, d) != h
    assert binding_hash("wezwanie_do_zaplaty", "0.1.0", f, {"uzasadnienie": d["uzasadnienie"] + " "}) != h
    assert binding_hash("x", "1", {}, None) == binding_hash("x", "1", {}, {})
    # NFC vs NFD of the same text are the same fact
    assert binding_hash("x", "1", {"a": "zażółć"}, None) == binding_hash("x", "1", {"a": "zażółć"}, None)


# --------------------------------------------------------------------------- complete drafts


@pytest.mark.parametrize("tid", TEMPLATE_IDS)
def test_complete_draft_exports_letter_and_separate_report(store, tmp_path, tid):
    f, d = facts_of(tid), draft_of(tid)
    rid = gen.synthetic_report(store, tid, f, d)
    r = render(store, tid, f, d, rid, tmp_path / "out")
    assert r.status == ResultStatus.ok, r.data
    assert r.data["document_status"] == DOCUMENT_COMPLETE
    assert r.data["report_id"] == rid and not r.data["placeholders"]
    files = {k: Path(v) for k, v in r.data["files"].items()}
    assert all(p.exists() and p.stat().st_size > 0 for p in files.values())

    letter_md = files["letter_md"].read_text(encoding="utf-8")
    letter_docx = docx_text(files["letter_docx"])
    for text in (letter_md, letter_docx):
        assert SIGNATURE_LINE in text
        assert "[UZUPEŁNIJ" not in text
        for banned in ("raport", "AI", "LLM", "eksperymentaln", "binding", "report_id", "snapshot",
                       "prawnie bezbłędn", "UZUPEŁNIJ", "get_legal_document"):
            assert banned not in text, banned
        assert "Jan Przykładowy" in text

    report_md = files["report_md"].read_text(encoding="utf-8")
    report_docx = docx_text(files["report_docx"]) + "\n" + "\n".join(
        c.text for t in Document(files["report_docx"]).tables for row in t.rows for c in row.cells)
    for text in (report_md, report_docx):
        assert STATUS_LINE in text
        assert rid in text
        assert "verified_exact" in text and "consolidated_text" in text
        assert "to NIE jest weryfikacja przez prawnika" in text  # reviewer_type=llm
        assert "prawnie bezbłędn" not in text.lower()
        assert gen.EXAMPLE_REPORT_NOTE in text
        assert "Wykorzystane fakty" in text and "Nierozwiązane problemy" in text


def test_draft_rendered_verbatim_only_in_allowed_section(store, tmp_path):
    tid = "wezwanie_do_zaplaty"
    f, d = facts_of(tid), {"uzasadnienie": "Pierwszy akapit.\n\nDrugi akapit – zażółć gęślą jaźń."}
    rid = gen.synthetic_report(store, tid, f, d)
    r = render(store, tid, f, d, rid, tmp_path)
    md = Path(r.data["files"]["letter_md"]).read_text(encoding="utf-8")
    assert "Pierwszy akapit." in md and "Drugi akapit – zażółć gęślą jaźń." in md
    bad = render(store, tid, f, {"uwagi_ai": "notatka"}, rid, tmp_path)
    assert bad.status == ResultStatus.invalid_input
    bad2 = render(store, "odstapienie_od_umowy_na_odleglosc", facts_of("odstapienie_od_umowy_na_odleglosc"),
                  {"uzasadnienie": "x"}, None, tmp_path)
    assert bad2.status == ResultStatus.invalid_input
    bad3 = render(store, tid, f, {"uzasadnienie": "To pismo jest prawnie bezbłędne."}, rid, tmp_path)
    assert bad3.status == ResultStatus.invalid_input


def test_interest_is_only_users_statement(store, tmp_path):
    tid = "wezwanie_do_zaplaty"
    f, d = facts_of(tid), draft_of(tid)
    rid = gen.synthetic_report(store, tid, f, d)
    md = Path(render(store, tid, f, d, rid, tmp_path).data["files"]["letter_md"]).read_text(encoding="utf-8")
    assert "wraz z odsetkami ustawowymi za opóźnienie od dnia 30.07.2026 r. do dnia zapłaty" in md
    amounts = re.findall(r"\d[\d ]*,\d{2} PLN", md)
    assert amounts == ["1 234,50 PLN"]  # no computed interest amount
    f2 = copy.deepcopy(f)
    f2["zadanie_odsetek"] = "nie"
    del f2["odsetki_od_dnia"]
    rid2 = gen.synthetic_report(store, tid, f2, d)
    md2 = Path(render(store, tid, f2, d, rid2, tmp_path).data["files"]["letter_md"]).read_text(encoding="utf-8")
    assert "odsetk" not in md2
    f3 = copy.deepcopy(f)
    del f3["odsetki_od_dnia"]  # interest claimed but start date missing -> unfilled form
    r3 = render(store, tid, f3, None, None, tmp_path)
    assert r3.data["document_status"] == DOCUMENT_INCOMPLETE
    assert any("odsetek" in p for p in r3.data["placeholders"])


def test_withdrawal_never_states_deadline_met(store, tmp_path):
    tid = "odstapienie_od_umowy_na_odleglosc"
    f = facts_of(tid)
    rid = gen.synthetic_report(store, tid, f, {})
    r = render(store, tid, f, {}, rid, tmp_path)
    md = Path(r.data["files"]["letter_md"]).read_text(encoding="utf-8").lower()
    for phrase in ("w terminie", "termin został", "zachowan", "przed upływem", "14 dni"):
        assert phrase not in md
    f["kwalifikacja"]["termin_wedlug_uzytkownika"] = "nie wiem"
    rid = gen.synthetic_report(store, tid, f, {})
    r = render(store, tid, f, {}, rid, tmp_path)
    assert r.status == ResultStatus.ok
    assert any("Zachowanie terminu nie jest potwierdzone" in w for w in r.warnings)


# --------------------------------------------------------------------------- export gates


def test_blocked_without_report(store, tmp_path):
    tid = "wezwanie_do_zaplaty"
    r = render(store, tid, facts_of(tid), draft_of(tid), None, tmp_path / "o")
    assert r.status == ResultStatus.blocked
    assert not (tmp_path / "o").exists() or not any((tmp_path / "o").iterdir())
    r = render(store, tid, facts_of(tid), draft_of(tid), "nie-istnieje", tmp_path / "o")
    assert r.status == ResultStatus.blocked and "Nie znaleziono" in r.data["reasons"][0]


def test_blocked_when_facts_draft_or_version_change(store, tmp_path):
    tid = "wezwanie_do_zaplaty"
    f, d = facts_of(tid), draft_of(tid)
    rid = gen.synthetic_report(store, tid, f, d)
    f2 = copy.deepcopy(f)
    f2["kwota"]["amount"] = "9999.00"
    del f2["kwota_slownie"]
    r = render(store, tid, f2, d, rid, tmp_path)
    assert r.status == ResultStatus.blocked and any("binding_hash" in x for x in r.data["reasons"])
    r = render(store, tid, f, {"uzasadnienie": d["uzasadnienie"] + " Dopisek."}, rid, tmp_path)
    assert r.status == ResultStatus.blocked
    rid_old = gen.synthetic_report(store, tid, f, d, binding_hash=binding_hash(tid, "0.0.9", f, d), report_id="stara-wersja")
    r = render(store, tid, f, d, rid_old, tmp_path)
    assert r.status == ResultStatus.blocked


def test_blocked_on_critical_errors_and_unverified_claims(store, tmp_path):
    tid = "reklamacja_konsumencka"
    f, d = facts_of(tid), draft_of(tid)
    rid = gen.synthetic_report(store, tid, f, d, critical_errors=["zmyślona sygnatura"], report_id="krytyczny")
    r = render(store, tid, f, d, rid, tmp_path)
    assert r.status == ResultStatus.blocked and any("krytyczne" in x for x in r.data["reasons"])
    for st in (CitationStatus.mismatch, CitationStatus.wrong_locator, CitationStatus.metadata_mismatch,
               CitationStatus.unmapped, CitationStatus.document_not_found, CitationStatus.version_mismatch):
        claims = [ClaimCheck(claim_id="c1", citation_status=st, temporal_status=TemporalStatus.confirmed)]
        rid = gen.synthetic_report(store, tid, f, d, claims=claims, report_id=f"zly-{st.value}".replace("_", "-"))
        r = render(store, tid, f, d, rid, tmp_path)
        assert r.status == ResultStatus.blocked, st
        assert "c1" in r.data["reasons"][0]


def test_blocked_when_snapshot_disappears(store, tmp_path):
    tid = "reklamacja_konsumencka"
    f, d = facts_of(tid), draft_of(tid)
    rid = gen.synthetic_report(store, tid, f, d, snapshot_ids=["saos:deadbeefdeadbeef"], report_id="bez-snapshotu")
    r = render(store, tid, f, d, rid, tmp_path)
    assert r.status == ResultStatus.blocked and any("Snapshoty" in x for x in r.data["reasons"])


def test_blocked_when_draft_text_has_no_checked_claims(store, tmp_path):
    tid = "wezwanie_do_zaplaty"
    f, d = facts_of(tid), draft_of(tid)
    rid = gen.synthetic_report(store, tid, f, d, claims=[], report_id="bez-twierdzen")
    assert render(store, tid, f, d, rid, tmp_path).status == ResultStatus.blocked
    # a withdrawal without free text may pass with an empty (but bound) report
    t2 = "odstapienie_od_umowy_na_odleglosc"
    rid2 = gen.synthetic_report(store, t2, facts_of(t2), {}, claims=[], report_id="odst-bez-twierdzen")
    assert render(store, t2, facts_of(t2), {}, rid2, tmp_path).status == ResultStatus.ok


def test_llm_review_never_presented_as_lawyer_and_temporal_unknown_listed(store, tmp_path):
    tid = "reklamacja_konsumencka"
    f, d = facts_of(tid), draft_of(tid)
    claims = [ClaimCheck(claim_id="c1", citation_status=CitationStatus.verified_exact,
                         temporal_status=TemporalStatus.unknown, reviewer_type=ReviewerType.llm)]
    rid = gen.synthetic_report(store, tid, f, d, claims=claims, report_id="czas-nieznany")
    r = render(store, tid, f, d, rid, tmp_path)
    assert r.status == ResultStatus.ok
    assert any("temporal_status=unknown" in w for w in r.warnings)
    rep = Path(r.data["files"]["report_md"]).read_text(encoding="utf-8")
    assert "NIE USTALONO brzmienia" in rep and "c1: nie ustalono brzmienia" in rep
    assert "sprawdzone przez prawnika" not in rep.replace("nie sprawdzony przez prawnika", "")


# --------------------------------------------------------------------------- out of scope / flags


@pytest.mark.parametrize("tid,qid,answer", [
    ("reklamacja_konsumencka", "konsument", "nie"),
    ("reklamacja_konsumencka", "przedmiot", "treść cyfrowa"),
    ("reklamacja_konsumencka", "przedmiot", "usługa cyfrowa"),
    ("reklamacja_konsumencka", "sprzedawca_przedsiebiorca", "nie"),
    ("odstapienie_od_umowy_na_odleglosc", "wyl_na_zamowienie", "tak"),
    ("odstapienie_od_umowy_na_odleglosc", "wyl_zapieczetowany_higiena", "tak"),
    ("odstapienie_od_umowy_na_odleglosc", "w_lokalu", "tak"),
    ("odstapienie_od_umowy_na_odleglosc", "konsument", "nie"),
    ("wezwanie_do_zaplaty", "wymagalnosc", "nie"),
    ("wezwanie_do_zaplaty", "postepowanie_w_toku", "tak"),
])
def test_out_of_scope(store, tmp_path, tid, qid, answer):
    f = facts_of(tid)
    f["kwalifikacja"][qid] = answer
    r = render(store, tid, f, draft_of(tid), None, tmp_path / "o")
    assert r.status == ResultStatus.out_of_scope
    assert r.data["reasons"]
    assert not (tmp_path / "o").exists()


def test_reklamacja_old_regime_flags_temporal_check_without_asserting(store, tmp_path):
    tid = "reklamacja_konsumencka"
    f, d = facts_of(tid), draft_of(tid)
    f["kwalifikacja"]["umowa_przed_zmiana_przepisow"] = "tak"
    rid = gen.synthetic_report(store, tid, f, d)
    r = render(store, tid, f, d, rid, tmp_path)
    assert r.status == ResultStatus.ok
    assert any("Wymagana kontrola czasowa" in w for w in r.warnings)
    letter = Path(r.data["files"]["letter_md"]).read_text(encoding="utf-8").lower()
    for word in ("rękojm", "ustaw", "art."):
        assert word not in letter  # letter does not assert which regime applies
    rep = Path(r.data["files"]["report_md"]).read_text(encoding="utf-8")
    assert "szablon nie przesądza" in rep.lower()


# --------------------------------------------------------------------------- incomplete form


@pytest.mark.parametrize("tid", TEMPLATE_IDS)
def test_incomplete_form_has_only_placeholders(store, tmp_path, tid):
    f = gen.incomplete_facts(tid)
    r = render(store, tid, f, draft_of(tid), "ignorowany", tmp_path)
    assert r.status == ResultStatus.ok
    assert r.data["document_status"] == DOCUMENT_INCOMPLETE
    assert r.data["report_id"] is None
    assert any("report_id pominięto" in w for w in r.warnings)
    md = Path(r.data["files"]["letter_md"]).read_text(encoding="utf-8")
    spec = load_template(tid)
    for name in gen.EXAMPLES[tid]["incomplete_drop"]:
        fs = spec.field(name)
        if fs.required or name == "data_odebrania_towaru":
            assert f"[UZUPEŁNIJ: {fs.label}]" in md
        orig = gen.EXAMPLES[tid]["facts"][name]
        if isinstance(orig, str) and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", orig):
            assert orig not in md  # nothing invented / leaked
    for draft_text in draft_of(tid).values():
        assert draft_text not in md  # unverified free text is not exported in an unfilled form
    rep = Path(r.data["files"]["report_md"]).read_text(encoding="utf-8")
    assert DOCUMENT_INCOMPLETE in rep and "Brakujące dane" in rep


def test_incomplete_when_qualifying_unanswered(store, tmp_path):
    tid = "wezwanie_do_zaplaty"
    f = facts_of(tid)
    del f["kwalifikacja"]["wymagalnosc"]
    r = render(store, tid, f, None, None, tmp_path)
    assert r.status == ResultStatus.ok and r.data["document_status"] == DOCUMENT_INCOMPLETE
    assert r.data["unanswered_questions"] == ["wymagalnosc"]


def test_empty_facts_give_blank_form_without_invented_values(store, tmp_path):
    tid = "wezwanie_do_zaplaty"
    r = render(store, tid, {}, None, None, tmp_path)
    assert r.status == ResultStatus.ok and r.data["document_status"] == DOCUMENT_INCOMPLETE
    md = Path(r.data["files"]["letter_md"]).read_text(encoding="utf-8")
    assert "[UZUPEŁNIJ: adres dłużnika]" in md
    assert not re.search(r"\d", md.replace("[UZUPEŁNIJ", ""))  # no invented dates/amounts/accounts
    assert "Załączniki" not in md


# --------------------------------------------------------------------------- validation


@pytest.mark.parametrize("tid,mutate,field", [
    ("wezwanie_do_zaplaty", lambda f: f["kwota"].update(amount="-5"), "kwota"),
    ("wezwanie_do_zaplaty", lambda f: f["kwota"].update(amount="0"), "kwota"),
    ("wezwanie_do_zaplaty", lambda f: f["kwota"].update(amount="10.555"), "kwota"),
    ("wezwanie_do_zaplaty", lambda f: f["kwota"].update(currency="XYZ"), "kwota"),
    ("wezwanie_do_zaplaty", lambda f: f.update(kwota=1234.5), "kwota"),
    ("wezwanie_do_zaplaty", lambda f: f.update(kwota_slownie="dwa tysiące złotych"), "kwota_slownie"),
    ("wezwanie_do_zaplaty", lambda f: f.update(numer_rachunku="12 3456 7890 1234 5678 9012 3456"), "numer_rachunku"),
    ("wezwanie_do_zaplaty", lambda f: f.update(termin_zaplaty="2026-09-01"), "termin_zaplaty"),
    ("wezwanie_do_zaplaty", lambda f: f.update(termin_wymagalnosci="2026-09-25"), "termin_wymagalnosci"),
    ("wezwanie_do_zaplaty", lambda f: f.update(odsetki_od_dnia="2026-07-01"), "odsetki_od_dnia"),
    ("wezwanie_do_zaplaty", lambda f: f.update(zadanie_odsetek="może"), "zadanie_odsetek"),
    ("wezwanie_do_zaplaty", lambda f: f.update(data_pisma="20.09.2026"), "data_pisma"),
    ("wezwanie_do_zaplaty", lambda f: f.update(data_pisma="2026-02-30"), "data_pisma"),
    ("wezwanie_do_zaplaty", lambda f: f.update(dluznik_adres="[UZUPEŁNIJ: adres dłużnika]"), "dluznik_adres"),
    ("reklamacja_konsumencka", lambda f: f.update(data_dostarczenia="2026-08-01"), "data_dostarczenia"),
    ("reklamacja_konsumencka", lambda f: f.update(data_zakupu="2026-12-01", data_dostarczenia="2026-12-02",
                                                  data_stwierdzenia_wady="2026-12-03", data_pisma="2026-12-04"), "data_zakupu"),
    ("reklamacja_konsumencka", lambda f: f.update(kupujacy_email="nie-email"), "kupujacy_email"),
    ("reklamacja_konsumencka", lambda f: f.update(zalaczniki="paragon"), "zalaczniki"),
    ("odstapienie_od_umowy_na_odleglosc", lambda f: f.update(data_odebrania_towaru="2026-09-01"), "data_odebrania_towaru"),
])
def test_validation_errors(store, tmp_path, tid, mutate, field):
    f = facts_of(tid)
    mutate(f)
    r = render(store, tid, f, None, None, tmp_path)
    assert r.status == ResultStatus.invalid_input, r.data
    assert field in {e["field"] for e in r.data["errors"]}


def test_invalid_qualifying_answer_and_unknown_template(store, tmp_path):
    f = facts_of("wezwanie_do_zaplaty")
    f["kwalifikacja"]["wymagalnosc"] = "raczej tak"
    assert render(store, "wezwanie_do_zaplaty", f, None, None, tmp_path).status == ResultStatus.invalid_input
    assert render(store, "nieznany", {}, None, None, tmp_path).status == ResultStatus.invalid_input


# --------------------------------------------------------------------------- DOCX structure


def test_docx_structure_diacritics_and_layout(store, tmp_path):
    tid = "reklamacja_konsumencka"
    f, d = facts_of(tid), draft_of(tid)
    f["opis_wady"] = "Zażółć gęślą jaźń – ZAŻÓŁĆ GĘŚLĄ JAŹŃ; wada obudowy."
    rid = gen.synthetic_report(store, tid, f, d)
    r = render(store, tid, f, d, rid, tmp_path)
    doc = Document(r.data["files"]["letter_docx"])
    paras = doc.paragraphs
    assert paras[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT and paras[0].text.startswith("Przykładowo, dnia")
    headings = [p for p in paras if p.style.name == "Heading 1"]
    assert [h.text for h in headings] == ["REKLAMACJA TOWARU"]
    assert headings[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert any("Zażółć gęślą jaźń – ZAŻÓŁĆ GĘŚLĄ JAŹŃ" in p.text for p in paras)
    numbered = [p.text for p in paras if p.style.name == "List Number"]
    assert numbered == ["kopia paragonu nr 0001/2026", "zdjęcie tabliczki znamionowej"]
    assert paras[-1].text == SIGNATURE_LINE and paras[-1].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    lang = doc.styles["Normal"].element.rPr.find(qn("w:lang"))
    assert lang.get(qn("w:val")) == "pl-PL"
    assert doc.core_properties.author == ""
    rep = Document(r.data["files"]["report_docx"])
    assert rep.paragraphs[0].style.name == "Heading 1"
    assert len(rep.tables) >= 3


# --------------------------------------------------------------------------- examples


@pytest.mark.parametrize("tid", TEMPLATE_IDS)
def test_examples_present_and_labelled(tid):
    for variant in ("kompletny", "niekompletny"):
        for part in ("pismo", "raport"):
            for ext in ("md", "docx"):
                assert (EXAMPLES_DIR / f"{tid}-{variant}-{part}.{ext}").exists()
        rep = (EXAMPLES_DIR / f"{tid}-{variant}-raport.md").read_text(encoding="utf-8")
        assert "DANE PRZYKŁADOWE" in rep and STATUS_LINE in rep
    rep = (EXAMPLES_DIR / f"{tid}-kompletny-raport.md").read_text(encoding="utf-8")
    assert gen.EXAMPLE_REPORT_NOTE in rep
    assert "[UZUPEŁNIJ" in (EXAMPLES_DIR / f"{tid}-niekompletny-pismo.md").read_text(encoding="utf-8")


def test_examples_regenerate_identically(tmp_path):
    res = gen.generate(tmp_path)
    for tid, r in res.items():
        assert r["complete"].status == ResultStatus.ok and r["complete"].data["document_status"] == DOCUMENT_COMPLETE
        assert r["incomplete"].data["document_status"] == DOCUMENT_INCOMPLETE
        for variant in ("kompletny", "niekompletny"):
            new = (tmp_path / f"{tid}-{variant}-pismo.md").read_text(encoding="utf-8")
            assert new == (EXAMPLES_DIR / f"{tid}-{variant}-pismo.md").read_text(encoding="utf-8")


@pytest.mark.skipif(shutil.which("qlmanage") is None, reason="qlmanage (macOS Quick Look) not available: page image not verified")
def test_docx_page_image_renders(tmp_path):
    src = EXAMPLES_DIR / "wezwanie_do_zaplaty-kompletny-pismo.docx"
    subprocess.run(["qlmanage", "-t", "-s", "800", "-o", str(tmp_path), str(src)],
                   capture_output=True, timeout=60, check=False)
    png = tmp_path / (src.name + ".png")
    assert png.exists() and png.stat().st_size > 10_000
