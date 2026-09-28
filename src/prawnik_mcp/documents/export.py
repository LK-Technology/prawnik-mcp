"""Markdown and DOCX writers for the letter and the separate sources/notes report.

The letter contains only the letter. All AI/verification notes go to the report.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from prawnik_mcp.contracts import CitationReport, CitationStatus, ReviewerType, SemanticReviewStatus, TemporalStatus
from prawnik_mcp.documents.schema import QualifyingQuestion, TemplateSpec

FONT = "Times New Roman"

CITATION_LABELS = {
    CitationStatus.verified_exact: "cytat zgodny dosłownie ze wskazaną wersją",
    CitationStatus.verified_normalized: "cytat zgodny po normalizacji odstępów i łączników",
    CitationStatus.mismatch: "cytat nie występuje we wskazanym miejscu",
    CitationStatus.wrong_locator: "cytat występuje pod innym lokalizatorem",
    CitationStatus.version_mismatch: "cytat tylko w innej wersji tekstu",
    CitationStatus.document_not_found: "dokumentu nie ma w lokalnym korpusie",
    CitationStatus.source_unavailable: "źródło niedostępne",
    CitationStatus.unmapped: "twierdzenie bez dowodu",
    CitationStatus.metadata_mismatch: "niezgodne metadane (np. sąd lub data)",
}
TEMPORAL_LABELS = {
    TemporalStatus.confirmed: "potwierdzony przedział obowiązywania",
    TemporalStatus.consolidated_text: "tekst jednolity (znana data stanu prawnego)",
    TemporalStatus.original_publication: "tekst w brzmieniu pierwotnym",
    TemporalStatus.unknown: "NIE USTALONO brzmienia dla daty sprawy",
    TemporalStatus.not_requested: "nie podano daty sprawy",
}
SEMANTIC_LABELS = {
    SemanticReviewStatus.not_performed: "nie przeprowadzono",
    SemanticReviewStatus.client_reported_pass: "zgłoszona przez klienta jako pozytywna (serwer tego nie weryfikował)",
    SemanticReviewStatus.client_reported_issues: "klient zgłosił problemy (serwer tego nie weryfikował)",
}
REVIEWER_LABELS = {
    ReviewerType.none: "brak recenzenta",
    ReviewerType.llm: "model językowy (LLM) uruchomiony przez klienta – to NIE jest weryfikacja przez prawnika",
    ReviewerType.human_unverified: "osoba zadeklarowana przez użytkownika – tożsamość i kompetencje niezweryfikowane",
}


@dataclass
class ReportContext:
    spec: TemplateSpec
    document_status: str
    binding_hash: str
    values: dict[str, Any]
    answers: dict[str, str]
    unanswered: list[QualifyingQuestion]
    missing: list[str]
    flags: list[str]
    warnings: list[str]
    report: CitationReport | None
    report_id: str | None
    example: bool
    generated_at: datetime


# --------------------------------------------------------------------------- letter


def letter_markdown(blocks) -> str:
    out: list[str] = []
    for b in blocks:
        if b.kind == "heading":
            out.append("# " + " ".join(b.lines))
        elif b.kind == "subheading":
            out.append("**" + " ".join(b.lines) + "**")
        elif b.kind == "numbered":
            out.append("\n".join(f"{i}. {x}" for i, x in enumerate(b.lines, 1)))
        elif b.kind == "signature":
            out.append("\n" + b.lines[0])
        else:
            out.append("  \n".join(b.lines))
    return "\n\n".join(out) + "\n"


def _base_document(title: str) -> Document:
    doc = Document()
    cp = doc.core_properties
    cp.title, cp.author, cp.last_modified_by, cp.comments = title, "", "", ""
    cp.language = "pl-PL"
    for style_name in ("Normal", "Heading 1", "Heading 2", "Title", "List Number", "List Bullet"):
        st = doc.styles[style_name]
        st.font.name = FONT
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), FONT)
        for t in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if fonts.get(qn(t)) is not None:
                del fonts.attrib[qn(t)]
        lang = rpr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            rpr.append(lang)
        lang.set(qn("w:val"), "pl-PL")
        if style_name.startswith("Heading") or style_name == "Title":
            st.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(11)
    return doc


def _lines_paragraph(doc, lines: list[str], align=None, bold=False):
    p = doc.add_paragraph()
    for i, line in enumerate(lines):
        run = p.add_run(line)
        run.bold = bold
        if i < len(lines) - 1:
            run.add_break()
    if align is not None:
        p.alignment = align
    return p


def letter_docx(blocks, title: str, path: Path) -> None:
    doc = _base_document(title)
    for b in blocks:
        if b.kind == "right":
            _lines_paragraph(doc, b.lines, WD_ALIGN_PARAGRAPH.RIGHT)
        elif b.kind == "left":
            _lines_paragraph(doc, b.lines, WD_ALIGN_PARAGRAPH.LEFT)
        elif b.kind == "heading":
            h = doc.add_heading(" ".join(b.lines), level=1)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            h.paragraph_format.space_before = Pt(18)
            h.paragraph_format.space_after = Pt(12)
        elif b.kind == "subheading":
            _lines_paragraph(doc, b.lines, bold=True)
        elif b.kind == "numbered":
            for item in b.lines:
                doc.add_paragraph(item, style="List Number")
        elif b.kind == "signature":
            p = _lines_paragraph(doc, b.lines, WD_ALIGN_PARAGRAPH.RIGHT)
            p.paragraph_format.space_before = Pt(36)
        else:
            _lines_paragraph(doc, b.lines, WD_ALIGN_PARAGRAPH.JUSTIFY)
    doc.save(path)


# --------------------------------------------------------------------------- report


def _fmt(spec: TemplateSpec, name: str, v: Any) -> str:
    from prawnik_mcp.documents.render import _display  # shared display rules

    return _display(spec.field(name), v)


def report_structure(ctx: ReportContext) -> list[tuple]:
    """Intermediate structure shared by the Markdown and DOCX writers."""
    from prawnik_mcp.documents.render import STATUS_LINE

    s = ctx.spec
    out: list[tuple] = [("h1", f"Raport źródeł i uwag – {s.title}")]
    out.append(("p", f"Status: {STATUS_LINE}."))
    out.append(("p", "Ten raport jest przeznaczony dla użytkownika i nie jest częścią pisma do adresata."))
    if ctx.example:
        out.append(("p", "DANE PRZYKŁADOWE: strony i dane w tym dokumencie są fikcyjne (np. „Jan Przykładowy”) – wyłącznie do demonstracji."))
    out.append(("h2", "Dokument"))
    out.append(("table", ["Pole", "Wartość"], [
        ["Szablon", f"{s.template_id} (wersja {s.version})"],
        ["Status szablonu", s.status],
        ["Status dokumentu", ctx.document_status],
        ["Powiązanie (binding_hash)", ctx.binding_hash],
        ["report_id", ctx.report_id or "brak (formularz niekompletny – raport niewymagany)"],
        ["Wygenerowano", ctx.generated_at.strftime("%Y-%m-%d %H:%M UTC")],
        ["Historia przeglądów szablonu", "brak przeglądu przez prawnika" if not s.review_history else f"{len(s.review_history)} wpis(y)"],
    ]))

    out.append(("h2", "Wykorzystane fakty"))
    rows = [[f.label, _fmt(s, f.name, ctx.values[f.name])] for f in s.required_facts if f.name in ctx.values]
    out.append(("table", ["Fakt", "Wartość podana przez użytkownika"], rows) if rows else ("p", "Brak podanych faktów."))
    if ctx.missing:
        out.append(("p", "Brakujące dane (w piśmie oznaczone jako [UZUPEŁNIJ: …]):"))
        out.append(("bullets", ctx.missing))

    out.append(("h2", "Pytania kwalifikujące"))
    qrows = [[qq.question, ctx.answers.get(qq.id, "BRAK ODPOWIEDZI")] for qq in s.qualifying_questions]
    out.append(("table", ["Pytanie", "Odpowiedź"], qrows))

    out.append(("h2", "Kontrola cytatów"))
    rep = ctx.report
    if rep is None:
        out.append(("p", "Nie przeprowadzono kontroli cytatów dla tego dokumentu (formularz niekompletny). Po uzupełnieniu danych wypełniony projekt wymaga raportu z check_citations."))
    else:
        out.append(("p", f"Raport: {rep.report_id}, utworzony {rep.created_at.isoformat()}. {rep.note}"))
        if rep.client_review_provided:
            out.append(("p", "Statusy recenzji semantycznej pochodzą od klienta AI, a nie z niezależnej kontroli serwera."))
        if rep.relevant_date:
            out.append(("p", f"Data sprawy przyjęta do kontroli czasowej: {rep.relevant_date.isoformat()}."))
        crow = [[c.claim_id, CITATION_LABELS.get(c.citation_status, c.citation_status.value) + f" ({c.citation_status.value})",
                 TEMPORAL_LABELS.get(c.temporal_status, c.temporal_status.value) + f" ({c.temporal_status.value})",
                 SEMANTIC_LABELS.get(c.semantic_review_status, c.semantic_review_status.value),
                 REVIEWER_LABELS.get(c.reviewer_type, c.reviewer_type.value)] for c in rep.claims]
        if crow:
            out.append(("table", ["Twierdzenie", "Status cytatu", "Status czasowy", "Recenzja semantyczna", "Recenzent"], crow))
        else:
            out.append(("p", "Raport nie zawiera twierdzeń."))
        if rep.snapshot_ids:
            out.append(("p", "Snapshoty źródeł, na których oparto kontrolę:"))
            out.append(("bullets", rep.snapshot_ids))

    out.append(("h2", "Kontrola czasowa"))
    tc = [f"{s.field(t['field']).label if s.field(t['field']) else t['field']}: {t['note']}" for t in s.temporal_checks]
    out.append(("bullets", tc + [s.supported_dates]))

    out.append(("h2", "Nierozwiązane problemy i uwagi"))
    issues: list[str] = []
    issues += [f"Brak odpowiedzi na pytanie kwalifikujące: {qq.question}" for qq in ctx.unanswered]
    issues += ctx.flags
    if rep is not None:
        issues += [f"Błąd krytyczny: {e}" for e in rep.critical_errors]
        for c in rep.claims:
            issues += [f"{c.claim_id}: {u}" for u in c.unresolved_issues]
            if c.temporal_status == TemporalStatus.unknown:
                issues.append(f"{c.claim_id}: nie ustalono brzmienia przepisu dla daty sprawy.")
            if c.semantic_review_status == SemanticReviewStatus.not_performed:
                issues.append(f"{c.claim_id}: nie oceniono, czy źródło wspiera twierdzenie i czy ma zastosowanie do faktów.")
    issues += [w for w in ctx.warnings if w not in issues]
    out.append(("bullets", issues or ["Brak zgłoszonych problemów. Nie oznacza to, że pismo jest poprawne prawnie."]))

    out.append(("h2", "Źródła do samodzielnej weryfikacji"))
    out.append(("p", s.legal_sources_note))
    out.append(("bullets", [f"{r.document_id} {r.locator or ''} – {r.purpose}".replace("  ", " ") for r in s.legal_sources]))

    out.append(("h2", "Ograniczenia"))
    out.append(("bullets", [
        "Dokument jest projektem eksperymentalnym i nie jest poradą prawną; szablonu nie sprawdził prawnik.",
        "Serwer sprawdza istnienie źródeł i wierność cytatów; nie ocenia, czy źródło wspiera twierdzenie ani czy prawo ma zastosowanie.",
        "Szablon nie oblicza terminów ani odsetek i nie stwierdza zachowania terminów.",
        "Recenzja wykonana przez model językowy nie jest recenzją prawnika.",
    ]))
    return out


def report_markdown(ctx: ReportContext) -> str:
    out: list[str] = []
    for item in report_structure(ctx):
        kind = item[0]
        if kind == "h1":
            out.append(f"# {item[1]}")
        elif kind == "h2":
            out.append(f"## {item[1]}")
        elif kind == "p":
            out.append(item[1])
        elif kind == "bullets":
            out.append("\n".join(f"- {x}" for x in item[1]))
        elif kind == "table":
            head, rows = item[1], item[2]
            esc = lambda x: str(x).replace("|", "\\|").replace("\n", " ")  # noqa: E731
            lines = ["| " + " | ".join(head) + " |", "|" + "---|" * len(head)]
            lines += ["| " + " | ".join(esc(c) for c in r) + " |" for r in rows]
            out.append("\n".join(lines))
    return "\n\n".join(out) + "\n"


def report_docx(ctx: ReportContext, path: Path) -> None:
    doc = _base_document(f"Raport źródeł i uwag – {ctx.spec.title}")
    for item in report_structure(ctx):
        kind = item[0]
        if kind == "h1":
            doc.add_heading(item[1], level=1)
        elif kind == "h2":
            doc.add_heading(item[1], level=2)
        elif kind == "p":
            doc.add_paragraph(item[1])
        elif kind == "bullets":
            for x in item[1]:
                doc.add_paragraph(x, style="List Bullet")
        elif kind == "table":
            head, rows = item[1], item[2]
            t = doc.add_table(rows=1, cols=len(head))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, h in enumerate(head):
                cell = t.rows[0].cells[i]
                cell.text = ""
                cell.paragraphs[0].add_run(h).bold = True
            for r in rows:
                cells = t.add_row().cells
                for i, c in enumerate(r):
                    cells[i].text = str(c)
    doc.save(path)
